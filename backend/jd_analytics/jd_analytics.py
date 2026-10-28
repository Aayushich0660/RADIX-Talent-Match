"""
Role 1: JD Analytics
Put this in backend/jd_analytics/jd_analytics.py

pip install langchain langchain-groq pdfplumber python-docx pydantic fastapi python-multipart python-dotenv

.env:
    GROQ_API_KEY=...
    JD_LLM_MODEL=groq:openai/gpt-oss-20b  # format "provider:model"
"""
import json
import os
import re
import sys
import tempfile
from pathlib import Path
from typing import List, Optional

import pdfplumber
from docx import Document
from dotenv import load_dotenv
from fastapi import APIRouter, UploadFile, File, HTTPException
from langchain.chat_models import init_chat_model
from langchain_core.prompts import ChatPromptTemplate
from pydantic import BaseModel

from backend.shared.schemas import Skill, SkillList, VALID_CODES  # adjust import to your repo layout

load_dotenv()

CACHE_DIR = Path("data/sample_outputs")
CACHE_DIR.mkdir(parents=True, exist_ok=True)

# ---------------------------------------------------------------- 1. Text extraction

def extract_text(file_path: str) -> str:
    """Return raw text from a .pdf or .docx file."""
    suffix = Path(file_path).suffix.lower()
    if suffix == ".pdf":
        with pdfplumber.open(file_path) as pdf:
            return "\n".join((page.extract_text() or "") for page in pdf.pages)
    if suffix == ".docx":
        doc = Document(file_path)
        parts = [p.text for p in doc.paragraphs]
        # JDs often hide content in tables
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    raise ValueError(f"Unsupported file type: {suffix}")


def clean_text(text: str) -> str:
    text = re.sub(r"\r", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"^\s*(Page\s*)?\d+\s*(of\s*\d+)?\s*$", "", text, flags=re.M | re.I)
    return text.strip()

# ---------------------------------------------------------------- 2. LLM extraction

class LLMOutput(BaseModel):
    """What we ask the LLM for. File name and source_type are filled in by code, not the LLM."""
    company: Optional[str]
    role: Optional[str]
    skills: List[Skill]


SYSTEM_PROMPT = """You are an expert technical recruiter. Extract the skills a job description requires.

Map each skill to EXACTLY ONE of these category codes:
- COD   : general programming/coding ability, specific programming languages (Python, Java, C++)
- DSA   : data structures, algorithms, problem solving, complexity
- OOD   : object-oriented design, design patterns, low-level design
- APTI  : aptitude, logical/analytical/quantitative reasoning
- COMM  : communication, collaboration, stakeholder work, documentation, presentation
- AI    : machine learning, deep learning, statistics for ML, GenAI/LLMs, data science modelling
- CLOUD : AWS/Azure/GCP, containers, Kubernetes, DevOps, CI/CD, deployment
- SQL   : SQL, relational databases, querying, data analysis with databases
- SWE   : software engineering practices: testing, code review, version control, APIs, debugging, agile, maintenance, production support
- SYSD  : system design, scalability, distributed systems, architecture, performance
- NETW  : networking, protocols, HTTP/TCP, security of network
- OS    : operating systems, Linux, concurrency, memory, processes
- OTHER : only if nothing above fits (e.g. Tableau, Excel, domain knowledge)

Rules:
1. The JD is written in natural prose. Infer skills from responsibilities, e.g.
   "own services handling millions of requests" -> SYSD (scalability).
   Do not rely on keyword matching only.
2. Focus most on "Key Responsibilities" and "What We're Looking For" (or similar) sections.
   Capture each distinct capability explicitly stated in those sections, including
   concrete subskills in lists and examples. Do not collapse distinct items such as
   SQL queries and primary keys, command-line use and log monitoring, batch and
   real-time processing, IP addresses and client-server architecture, or latency
   and structured problem-solving into a single skill. Also capture explicit skills
   stated in the role overview and "Nice to Have" sections. Separate paired skills
   and listed concepts such as normalization vs. denormalization, regression vs.
   classification, overfitting, model-output checks, sampling, probability, and
   metric definition instead of hiding them inside a broad category. If the JD
   explicitly mentions system security or secure systems, extract security as a
   separate skill even when it is not described as network security.
3. Scan the ENTIRE JD, including every section, for every named programming language,
   tool, library, framework, platform, database, and analytics or BI product. Always
   extract each distinct named item as its own separate skill, even if it appears only
   once, in a preferred qualification, or outside the main responsibilities. Examples
   include Python, R, Excel, Power BI, Tableau, SQL, pandas, and Kubernetes. Never
   combine multiple named technologies into one skill; do not omit a named item because
   it seems less important. Map each item to its most relevant category.
4. skill_name: short, canonical (e.g. "Python", "REST APIs", "Distributed Systems").
5. evidence: a SHORT quote (under 25 words) copied from the JD, or a one-line reason.
6. confidence: "high" = explicitly required/must-have; "medium" = implied by duties; "low" = nice-to-have/preferred/bonus.
7. Deduplicate exact repeats and true synonyms (case-insensitive skill name), but do
   not combine distinct subskills just because they share a topic or appear in one
   sentence. Keep each separately named capability as its own skill. Do not invent
   unsupported skills.
8. Also return the company name and role title if present in the text (else null). Return
   the role as a clean job title only: remove any parenthetical or bracketed abbreviation,
   level, or qualifier (for example, "Software Engineer (SWE)" -> "Software Engineer").
   Do not include brackets or bracketed text in the role.
9. Be exhaustive. Go through every sentence of 'Key Responsibilities' and 'What We're Looking For' and extract every distinct skill. A typical JD yields 12 to 25 skills. Do not stop early."""

_prompt = ChatPromptTemplate.from_messages([
    ("system", SYSTEM_PROMPT),
    ("human", "Job description:\n\n{jd_text}"),
])

_llm = None

def _get_chain():
    global _llm
    if _llm is None:
        model = os.getenv("JD_LLM_MODEL", "openai:gpt-4o-mini")
        _llm = init_chat_model(
            model,
            temperature=0,
            max_tokens=6000,
            reasoning_effort="low",
            reasoning_format="hidden",
        )
    return _prompt | _llm.with_structured_output(
        LLMOutput,
        method="json_schema",
        strict=True,
    )

# ---------------------------------------------------------------- 3. Post-processing

ALIASES = {
    "py": "Python", "js": "JavaScript", "ts": "TypeScript", "postgres": "PostgreSQL",
    "ml": "Machine Learning", "k8s": "Kubernetes", "restful apis": "REST APIs",
    "rest api": "REST APIs", "dsa": "Data Structures & Algorithms",
}
CONF_RANK = {"high": 3, "medium": 2, "low": 1}


def postprocess(skills: List[Skill]) -> List[Skill]:
    best = {}
    for s in skills:
        name = ALIASES.get(s.skill_name.strip().lower(), s.skill_name.strip())
        code = s.category_code if s.category_code in VALID_CODES else "OTHER"
        key = name.lower()
        fixed = Skill(skill_name=name, category_code=code,
                      evidence=s.evidence.strip(), confidence=s.confidence)
        # on duplicates keep the higher-confidence entry
        if key not in best or CONF_RANK[fixed.confidence] > CONF_RANK[best[key].confidence]:
            best[key] = fixed
    return list(best.values())


def clean_role_title(role: Optional[str]) -> Optional[str]:
    if role is None:
        return None
    title = re.sub(r"\s*[\(\[].*?[\)\]]", "", role).strip()
    return title or None


def add_explicit_security_skill(text: str, skills: List[Skill]) -> List[Skill]:
    if any("security" in skill.skill_name.lower() for skill in skills):
        return skills
    sentence = next(
        (
            sentence.strip()
            for sentence in re.split(r"(?<=[.!?])\s+", text)
            if re.search(r"\b(?:secure|security)\b", sentence, re.IGNORECASE)
        ),
        None,
    )
    if sentence is not None:
        evidence = " ".join(sentence.split()[:24])
        skills.append(
            Skill(
                skill_name="Security",
                category_code="SYSD",
                evidence=evidence,
                confidence="medium",
            )
        )
    return skills


def extract_consistent_output(text: str) -> LLMOutput:
    chain = _get_chain()
    outputs = [chain.invoke({"jd_text": text}) for _ in range(2)]
    best_output = max(outputs, key=lambda output: len(output.skills))
    return LLMOutput(
        company=best_output.company,
        role=best_output.role,
        skills=postprocess(
            add_explicit_security_skill(text, postprocess(best_output.skills))
        ),
    )


# ---------------------------------------------------------------- 4. Main function

def analyze_jd(file_path: str, use_cache: bool = True) -> SkillList:
    name = Path(file_path).name
    cache_file = CACHE_DIR / f"jd_{Path(file_path).stem}.json"
    if use_cache and cache_file.exists():
        return SkillList(**json.loads(cache_file.read_text(encoding="utf-8")))

    text = clean_text(extract_text(file_path))
    if len(text) < 100:
        raise ValueError("Could not read enough text from the file (scanned PDF?)")

    out = extract_consistent_output(text)
    result = SkillList(
        source_type="jd",
        source_file=name,
        company=out.company,
        role=clean_role_title(out.role),
        skills=postprocess(out.skills),
    )
    cache_file.write_text(result.model_dump_json(indent=2), encoding="utf-8")
    return result

# ---------------------------------------------------------------- 5. API

router = APIRouter(prefix="/api")
JD_DIR = Path("data/jds")


@router.get("/jds")
def list_jds():
    """Pre-loaded JDs so the UI can offer 'pick a JD'."""
    return sorted(p.name for p in JD_DIR.glob("*") if p.suffix.lower() in {".pdf", ".docx"})


@router.post("/jd/analyze", response_model=SkillList)
async def analyze_upload(file: UploadFile = File(...)):
    suffix = Path(file.filename).suffix.lower()
    if suffix not in {".pdf", ".docx"}:
        raise HTTPException(400, "Upload a PDF or DOCX")
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp.write(await file.read())
        tmp_path = tmp.name
    try:
        result = analyze_jd(tmp_path, use_cache=False)
        result.source_file = file.filename
        return result
    except Exception as e:
        raise HTTPException(500, str(e))
    finally:
        os.unlink(tmp_path)


@router.post("/jd/analyze-sample/{filename}", response_model=SkillList)
def analyze_sample(filename: str):
    path = JD_DIR / filename
    if not path.exists():
        raise HTTPException(404, "JD not found")
    return analyze_jd(str(path))


# ---------------------------------------------------------------- CLI test
import time

def analyze_with_retry(path, tries=4):
    for attempt in range(tries):
        try:
            return analyze_jd(path, use_cache=False)
        except Exception as e:
            if "429" in str(e) or "rate" in str(e).lower():
                wait = 20 * (attempt + 1)
                print(f"Rate limit hit, waiting {wait} seconds...")
                time.sleep(wait)
            else:
                raise
    raise RuntimeError("Still rate limited, try again in a few minutes")


if __name__ == "__main__":
    sys.stdout.reconfigure(errors="replace")
    for f in sorted(JD_DIR.glob("*.pdf")):
        r = analyze_with_retry(str(f))
        print(f"\n=== {r.company} | {r.role} ({len(r.skills)} skills)")
        for s in r.skills:
            print(f"  [{s.category_code:5}] {s.skill_name:35} {s.confidence}")
        time.sleep(15)