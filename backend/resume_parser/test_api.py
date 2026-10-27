import io
import tempfile
import unittest
from pathlib import Path

from fastapi.testclient import TestClient



from api import MAX_FILE_SIZE_BYTES, app
from heuristic import heuristic_extract


def make_pdf() -> bytes:
    content = b"BT /F1 12 Tf 72 720 Td (Jordan Rivera) Tj 0 -18 Td (jordan.rivera@example.test) Tj 0 -18 Td (+1 415 890 2341) Tj 0 -18 Td (Education) Tj 0 -18 Td (Bachelors Computer Science Test University 2026) Tj 0 -18 Td (Skills Python TypeScript React) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode()
    )
    return bytes(output)


def make_docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph("Jordan Rivera")
    document.add_paragraph("jordan.rivera@example.test")
    document.add_paragraph("+1 415 890 2341")
    document.add_paragraph("Skills: Python, TypeScript, React")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


class ResumeParserApiTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def test_keyword_matching_normalizes_pdf_line_breaks_and_preserves_evidence(self):
        spaced = heuristic_extract("process scheduling").skills
        wrapped = heuristic_extract("process\nscheduling").skills

        process_scheduling = next(skill for skill in spaced if skill.skill_name == "Process Scheduling")
        wrapped_process_scheduling = next(
            (skill for skill in wrapped if skill.skill_name == "Process Scheduling"),
            None,
        )

        self.assertIsNotNone(wrapped_process_scheduling)
        self.assertEqual(wrapped_process_scheduling.model_dump(), process_scheduling.model_dump())
        self.assertEqual(wrapped_process_scheduling.evidence, "process scheduling")

    def test_keyword_matching_ignores_role_fit_mentions_but_keeps_technical_skills(self):
        role_fit = heuristic_extract(
            "Entry-level, support-oriented — best fit for the Oracle FSS roles, "
            "weaker fit for Google / Microsoft"
        )
        technical_skill = heuristic_extract("Technical Skills\nOracle Database")

        self.assertFalse(any(skill.skill_name == "Oracle" for skill in role_fit.skills))
        self.assertTrue(any(skill.skill_name == "Oracle" for skill in technical_skill.skills))

    def test_existing_priya_resume_extracts_projects_experience_hackathon_and_certification(self):
        fixtures = Path(__file__).parent / "test_resumes"
        for path in (
            fixtures / "PDF" / "Priya Menon.pdf",
            fixtures / "Word" / "Priya Menon.docx",
        ):
            with self.subTest(filename=path.name, extension=path.suffix):
                content_type = (
                    "application/pdf"
                    if path.suffix.lower() == ".pdf"
                    else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
                response = self.client.post(
                    "/parse-resume",
                    files={"file": (path.name, path.read_bytes(), content_type)},
                )
                self.assertEqual(response.status_code, 200, response.text)
                result = response.json()
                projects = {project["title"]: project for project in result["projects"]}
                self.assertIn("simple student attendance tracker", projects)
                self.assertIn("college helpdesk simulation project", projects)
                self.assertIn("12-hour campus hackathon", projects)
                self.assertEqual(
                    projects["12-hour campus hackathon"]["description"],
                    "Participated (non-finalist)",
                )
                self.assertEqual(
                    [
                        (item["role"], item["company"], item["description"])
                        for item in result["experience"]
                    ],
                    [(
                        "Summer intern, IT support desk",
                        "a local business",
                        "handled first-level troubleshooting and escalations",
                    )],
                )
                self.assertEqual(
                    result["certifications"],
                    ["Cisco Networking Basics (online course)"],
                )
                self.assertEqual(
                    result["preferred_roles"],
                    ["Application Support Analyst", "Associate Software Engineer"],
                )
                attendance_project = next(
                    project for project in result["projects"]
                    if project["title"] == "simple student attendance tracker"
                )
                self.assertEqual(attendance_project["technologies"], ["Java", "MySQL"])
                ticketing_skill = next(
                    skill for skill in result["skills"]
                    if skill["skill_name"] == "Ticketing Tools"
                )
                self.assertEqual(ticketing_skill["category_code"], "OTHER")
                self.assertEqual(ticketing_skill["evidence"], "exposure to ticketing tools")

    def test_existing_ananya_resume_returns_education_fields(self):
        path = Path(__file__).parent / "test_resumes" / "PDF" / "Ananya Rao.pdf"
        response = self.client.post(
            "/parse-resume",
            files={"file": (path.name, path.read_bytes(), "application/pdf")},
        )

        self.assertEqual(response.status_code, 200, response.text)
        education = response.json()["education"]
        self.assertEqual(len(education), 1)
        self.assertEqual(education[0]["degree"], "B.Tech")
        self.assertEqual(education[0]["field"], "Computer Science and Engineering")
        self.assertEqual(education[0]["institution"], "")
        self.assertEqual(education[0]["end_year"], "2026")
        self.assertEqual(education[0]["grade"], "CGPA 8.6 / 10")
        self.assertEqual(education[0]["cgpa"], "8.6/10")

    def test_parser_does_not_infer_cgpa_from_a_gpa_grade(self):
        result = heuristic_extract("Education\nB.Tech, Computer Science GPA 3.9/4.0")

        self.assertEqual(len(result.education), 1)
        self.assertEqual(result.education[0].grade, "GPA 3.9/4.0")
        self.assertIsNone(result.education[0].cgpa)

    def test_karthik_technical_skills_and_evidence_are_preserved_for_pdf_and_docx(self):
        fixtures = Path(__file__).parent / "test_resumes"
        expected = {
            "SQL": ("SQL", "SQL"),
            "Python": ("COD", "Python (pandas, scikit-learn)"),
            "Java": ("COD", "basic Java"),
            "Pandas": ("COD", "Python (pandas, scikit-learn)"),
            "Scikit-Learn": ("AI", "Python (pandas, scikit-learn)"),
            "Power BI": ("OTHER", "Power BI"),
            "Statistics": ("OTHER", "Statistics"),
            "A/B Testing": ("OTHER", "A/B testing basics"),
            "Excel": ("OTHER", "Excel"),
        }
        for path in (
            fixtures / "PDF" / "Karthik Subramaniam.pdf",
            fixtures / "Word" / "Karthik Subramaniam.docx",
        ):
            with self.subTest(filename=path.name, extension=path.suffix):
                content_type = (
                    "application/pdf"
                    if path.suffix.lower() == ".pdf"
                    else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
                response = self.client.post(
                    "/parse-resume",
                    files={"file": (path.name, path.read_bytes(), content_type)},
                )
                self.assertEqual(response.status_code, 200, response.text)
                skills = {
                    skill["skill_name"]: skill
                    for skill in response.json()["skills"]
                }
                for name, (category, evidence) in expected.items():
                    with self.subTest(skill=name):
                        self.assertIn(name, skills)
                        self.assertEqual(skills[name]["category_code"], category)
                        self.assertEqual(skills[name]["evidence"], evidence)

    def test_health_reports_offline_parser_default(self):
        response = self.client.get("/health")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"ok": True, "mode": "offline_heuristic"})

    def test_pdf_upload_uses_offline_parser_and_returns_role2_contract(self):
        response = self.client.post(
            "/parse-resume",
            files={"file": ("sample.pdf", make_pdf(), "application/pdf")},
        )
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body["source_type"], "resume")
        self.assertEqual(body["source_file"], "sample.pdf")
        self.assertEqual(body["candidate"]["name"], "Jordan Rivera")
        self.assertEqual(body["candidate"]["email"], "jordan.rivera@example.test")
        self.assertTrue(any(skill["skill_name"] == "Python" for skill in body["skills"]))
        self.assertEqual(set(body), {
            "source_type", "source_file", "company", "role", "candidate",
            "education", "projects", "experience", "certifications",
            "preferred_roles", "skills",
        })

    def test_docx_upload_is_read_and_extracted(self):
        response = self.client.post(
            "/parse-resume",
            files={"file": ("sample.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body["candidate"]["name"], "Jordan Rivera")
        self.assertTrue(any(skill["skill_name"].lower() == "typescript" for skill in body["skills"]))

    def test_rejects_unsupported_empty_oversized_and_malformed_files(self):
        unsupported = self.client.post(
            "/parse-resume",
            files={"file": ("resume.txt", b"resume", "text/plain")},
        )
        self.assertEqual(unsupported.status_code, 400)

        empty = self.client.post(
            "/parse-resume",
            files={"file": ("empty.pdf", b"", "application/pdf")},
        )
        self.assertEqual(empty.status_code, 422)

        oversized = self.client.post(
            "/parse-resume",
            files={"file": ("large.pdf", b"%PDF-" + b"x" * MAX_FILE_SIZE_BYTES, "application/pdf")},
        )
        self.assertEqual(oversized.status_code, 413)

        invalid_signature = self.client.post(
            "/parse-resume",
            files={"file": ("fake.pdf", b"not a PDF", "application/pdf")},
        )
        self.assertEqual(invalid_signature.status_code, 422)

        malformed_docx = self.client.post(
            "/parse-resume",
            files={"file": ("broken.docx", b"PK broken archive", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        self.assertEqual(malformed_docx.status_code, 422)

        malformed_pdf = self.client.post(
            "/parse-resume",
            files={"file": ("broken.pdf", b"%PDF-1.4 broken", "application/pdf")},
        )
        self.assertEqual(malformed_pdf.status_code, 422)

    def test_temporary_upload_file_is_removed_after_parsing(self):
        previous_tempdir = tempfile.tempdir
        try:
            with tempfile.TemporaryDirectory() as test_tempdir:
                tempfile.tempdir = test_tempdir
                response = self.client.post(
                    "/parse-resume",
                    files={"file": ("sample.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                )
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(list(Path(test_tempdir).iterdir()), [])
        finally:
            tempfile.tempdir = previous_tempdir


if __name__ == "__main__":
    unittest.main()
