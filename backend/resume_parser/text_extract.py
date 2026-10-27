"""Extract text from PDF and DOCX documents."""

import re
from pathlib import Path


def _clean(text: str) -> str:
    text = text.replace("\u00a0", " ").replace("\u2022", "- ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdf(path: Path) -> str:
    text = ""
    try:
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            text = "\n\n".join((page.extract_text() or "") for page in pdf.pages)
    except Exception:
        text = ""
    if len(text.strip()) < 50:
        from pypdf import PdfReader

        text = "\n\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)
    return text


def _docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            seen, cells = set(), []
            for cell in row.cells:
                if cell._tc not in seen:
                    seen.add(cell._tc)
                    cells.append(cell.text.strip())
            parts.append(" | ".join(value for value in cells if value))
    for section in document.sections:
        parts = [paragraph.text for paragraph in section.header.paragraphs] + parts
    return "\n".join(parts)


def extract_text(path: str) -> str:
    document_path = Path(path)
    extension = document_path.suffix.lower()
    try:
        if extension == ".pdf":
            raw = _pdf(document_path)
        elif extension == ".docx":
            raw = _docx(document_path)
        else:
            raise ValueError(f"Unsupported file type: {extension} (use PDF or DOCX)")
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("The document is malformed or cannot be read.") from error

    text = _clean(raw)
    if not text:
        raise ValueError("No text could be extracted (scanned PDF? try OCR).")
    return text
